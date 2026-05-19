"""Extract plain text from order content, history entries and S3-stored files."""
import json
import logging
from io import BytesIO

from botocore.exceptions import BotoCoreError, ClientError
from flask import current_app

from application.file_utils import get_s3_client, is_s3_storage_enabled, is_s3_storage_path

logger = logging.getLogger(__name__)


SUPPORTED_FILE_EXTENSIONS = {'pdf', 'docx', 'xlsx', 'txt'}


def _file_extension(filename: str) -> str:
    return filename.rsplit('.', 1)[1].lower() if '.' in filename else ''


def extract_order_content(order) -> str | None:
    """Build the text payload for an order's main content."""
    parts = []
    if order.title:
        parts.append(f'Заголовок: {order.title}')
    if order.priority:
        parts.append(f'Приоритет: {order.priority}')
    if order.status:
        parts.append(f'Статус: {order.status}')
    if order.creator_name:
        parts.append(f'Автор: {order.creator_name}')
    if order.deadline:
        parts.append(f'Срок: {order.deadline.isoformat()}')
    if order.content:
        parts.append('')
        parts.append(order.content)
    text = '\n'.join(parts).strip()
    return text or None


def extract_order_history(history_entries, order_result=None) -> str | None:
    """Concatenate order history details and the final result into one document."""
    lines = []
    for entry in history_entries:
        timestamp = entry.created_at.isoformat() if entry.created_at else ''
        actor = entry.user_name or ''
        action = entry.action or ''
        details = entry.details or ''
        lines.append(f'[{timestamp}] {actor} — {action}: {details}'.strip())
    if order_result:
        if isinstance(order_result, dict):
            content = order_result.get('content')
            if content:
                lines.append('')
                lines.append(f'Результат исполнения: {content}')
        else:
            lines.append('')
            lines.append(f'Результат исполнения: {json.dumps(order_result, ensure_ascii=False)}')
    text = '\n'.join(line for line in lines if line).strip()
    return text or None


def _download_file_bytes(order_file) -> bytes | None:
    if not is_s3_storage_enabled():
        logger.warning('S3 storage is not configured; cannot fetch file %s', order_file.id)
        return None
    if not is_s3_storage_path(order_file.filepath):
        logger.info('File %s is not stored in S3, skipping', order_file.id)
        return None
    try:
        response = get_s3_client().get_object(
            Bucket=current_app.config['S3_BUCKET_NAME'],
            Key=order_file.filepath,
        )
    except (BotoCoreError, ClientError) as exc:
        logger.warning('Failed to fetch %s from S3: %s', order_file.filepath, exc)
        return None
    return response['Body'].read()


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover
        logger.error('pypdf is not installed; PDF extraction disabled')
        return ''
    reader = PdfReader(BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or '')
        except Exception:  # pragma: no cover - defensive
            logger.exception('Failed to extract a PDF page')
    return '\n'.join(pages).strip()


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError:  # pragma: no cover
        logger.error('python-docx is not installed; DOCX extraction disabled')
        return ''
    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text]
            if cells:
                parts.append(' | '.join(cells))
    return '\n'.join(parts).strip()


def _extract_xlsx(data: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover
        logger.error('openpyxl is not installed; XLSX extraction disabled')
        return ''
    wb = load_workbook(BytesIO(data), data_only=True, read_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f'# Лист: {sheet.title}')
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(' | '.join(cells))
    return '\n'.join(parts).strip()


def _extract_txt(data: bytes) -> str:
    for encoding in ('utf-8', 'utf-8-sig', 'cp1251', 'latin-1'):
        try:
            return data.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return data.decode('utf-8', errors='ignore').strip()


def extract_file_text(order_file) -> str | None:
    """Extract searchable text from an attached order file in S3."""
    ext = _file_extension(order_file.original_name or order_file.filename or '')
    if ext not in SUPPORTED_FILE_EXTENSIONS:
        logger.info('Skipping unsupported file %s (extension=%s)', order_file.original_name, ext)
        return None
    data = _download_file_bytes(order_file)
    if not data:
        return None
    try:
        if ext == 'pdf':
            text = _extract_pdf(data)
        elif ext == 'docx':
            text = _extract_docx(data)
        elif ext == 'xlsx':
            text = _extract_xlsx(data)
        elif ext == 'txt':
            text = _extract_txt(data)
        else:
            text = ''
    except Exception:
        logger.exception('Failed to extract text from %s', order_file.original_name)
        return None
    text = (text or '').strip()
    return text or None
